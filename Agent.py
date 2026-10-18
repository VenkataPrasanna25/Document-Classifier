import os, re, json, shutil, time, threading
from collections import defaultdict
from pathlib import Path

from dotenv import load_dotenv
from openai import OpenAI
from pypdf import PdfReader
from docx import Document
from watchdog.observers import Observer
from watchdog.events import FileSystemEventHandler

load_dotenv()
client = OpenAI(api_key=os.getenv("LLM_API_KEY"),
                base_url=os.getenv("LLM_BASE_URL"))
MODEL = os.getenv("LLM_MODEL")
BASE=Path(__file__).resolve().parent
INBOX = BASE / "inbox"
OUT = BASE / "sorted"
MAX_CATEGORIES = 12
FALLBACK = "Unsorted"
OVERFLOW = "Other"
MAX_ATTEMPTS = 3

LOCK = threading.Lock()
attempts = defaultdict(int)

# ---------- OCR (optional) ----------
_ocr = None
def ocr_image(img) -> str:
    global _ocr
    try:
        if _ocr is None:
            RapidOCR = __import__("rapidocr_onnxruntime", fromlist=["RapidOCR"]).RapidOCR
            _ocr = RapidOCR()
        result, _ = _ocr(img)
        return "\n".join(r[1] for r in (result or []))
    except Exception as e:
        print(f"  (OCR unavailable: {e})")
        return ""

def ocr_pdf(path: Path) -> str:
    try:
        pdfium = __import__("pypdfium2")
        pdf = pdfium.PdfDocument(str(path))
        texts = []
        for i in range(min(2, len(pdf))):
            img = pdf[i].render(scale=2).to_numpy()
            texts.append(ocr_image(img))
        return "\n".join(texts)
    except Exception as e:
        print(f"  (PDF OCR unavailable: {e})")
        return ""

# ---------- Tool 1: read the document ----------
def extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    text = ""
    try:
        if ext == ".pdf":
            reader = PdfReader(path)
            text = "\n".join((p.extract_text() or "") for p in reader.pages[:3])
            if len(text.strip()) < 30:             # probably scanned
                print("  no text layer, trying OCR...")
                text = ocr_pdf(path)
        elif ext == ".docx":
            text = "\n".join(p.text for p in Document(path).paragraphs)
        elif ext in (".txt", ".md", ".csv", ".log", ".json", ".html", ".htm", ".xml"):
            text = path.read_text(errors="ignore")
        elif ext == ".xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=True)
            rows = []
            for ws in wb.worksheets[:2]:
                rows.append(f"[Sheet: {ws.title}]")
                for row in ws.iter_rows(max_row=40, values_only=True):
                    rows.append(" | ".join(str(c) for c in row if c is not None))
            text = "\n".join(rows)
        elif ext == ".pptx":
            from pptx import Presentation
            parts = []
            for slide in Presentation(path).slides:
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        parts.append(shape.text_frame.text)
            text = "\n".join(parts)
        elif ext in (".png", ".jpg", ".jpeg", ".bmp"):
            text = ocr_image(str(path))
    except Exception as e:
        print(f"  ! could not read {path.name}: {e}")
    return text.strip()

# ---------- Category memory = existing folders ----------
def existing_categories() -> list[str]:
    if not OUT.exists():
        return []
    return sorted(d.name for d in OUT.iterdir()
                  if d.is_dir() and d.name not in (FALLBACK, OVERFLOW))

def clean_name(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9 _-]", "", str(name)).strip().title()[:40]

def key(name: str) -> str:
    k = name.lower().replace(" ", "").replace("-", "").replace("_", "")
    return k[:-1] if k.endswith("s") else k

def match_existing(name: str, existing: list[str]):
    for e in existing:
        if key(e) == key(name):
            return e
    return None

# ---------- Brain: ask the LLM ----------
def ask_llm(prompt: str) -> str:
    for attempt in range(3):
        try:
            r = client.chat.completions.create(
                model=MODEL,
                messages=[{"role": "user", "content": prompt}],
                temperature=0,
            )
            content = r.choices[0].message.content or ""
            if not content.strip():
                print("  ! LLM returned an empty reply")
            else:
                return content
        except Exception as e:
            print(f"  ! LLM error: {e}")
        time.sleep(5 * (attempt + 1))
    return ""

def classify(filename: str, text: str):
    """Returns a category name, or None if the LLM failed (so we can retry)."""
    existing = existing_categories()
    can_create = len(existing) < MAX_CATEGORIES
    existing_txt = "\n".join(f"- {c}" for c in existing) or "(none yet)"
    rule = ("If none of the existing categories fit, create a NEW category."
            if can_create else
            "You may NOT create new categories. Pick the closest existing one.")
    if text:
        body = f'Document content (truncated):\n"""{text[:3000]}"""'
    else:
        body = ("Document content: NOT AVAILABLE. Decide using the file name "
                "and file extension only.")

    prompt = f"""You are a document organiser. Read the document and decide which folder it belongs in.

Existing categories:
{existing_txt}

Rules:
- Base your decision on the document CONTENT, not just the file name.
- Prefer reusing an existing category when the document fits it.
- {rule}
- New category names must be broad and general, 1-3 words, Title Case.
  Good: Invoices, Resumes, Legal Contracts, Research Papers, Meeting Notes.
  Bad: Invoice From Acme March 2024, Misc Stuff, Documents.

File name: {filename}
{body}

Reply ONLY with JSON, nothing else:
{{"category": "<name>", "reason": "<short>"}}"""

    content = ask_llm(prompt)
    if not content:
        return None
    try:
        blocks = re.findall(r"\{[^{}]*\}", content, re.S)
        data = json.loads(blocks[-1])
        name = clean_name(data.get("category", ""))
        print(f"  LLM suggested: '{name}' - {data.get('reason')}")
        if not name:
            return None
        match = match_existing(name, existing)
        if match:
            return match
        return name if can_create else OVERFLOW
    except Exception as e:
        print(f"  ! bad LLM output: {e}\n    raw: {content[:300]}")
        return None

# ---------- Tool 2: move the file ----------
def wait_until_stable(path: Path, tries=40) -> bool:
    last = -1
    for _ in range(tries):
        try:
            size = path.stat().st_size
        except FileNotFoundError:
            return False
        if size == last and size > 0:
            return True
        last = size
        time.sleep(0.5)
    return False

def unique_target(folder: Path, name: str) -> Path:
    target = folder / name
    i = 1
    while target.exists():
        target = folder / f"{Path(name).stem}_{i}{Path(name).suffix}"
        i += 1
    return target

def process(path: Path):
    with LOCK:
        if not path.is_file() or path.name.startswith("."):
            return
        print(f"\nNew file: {path.name}")
        if not wait_until_stable(path):
            return
        text = extract_text(path)
        print(f"  extracted {len(text)} characters")
        category = classify(path.name, text)
        if category is None:
            attempts[str(path)] += 1
            if attempts[str(path)] < MAX_ATTEMPTS:
                print(f"  LLM failed (attempt {attempts[str(path)]}/{MAX_ATTEMPTS}),"
                      " leaving file in inbox to retry")
                return
            category = FALLBACK
        dest_dir = OUT / category
        dest_dir.mkdir(parents=True, exist_ok=True)
        if not path.exists():
            print(f"  file already handled,skipping")
            return
        dest = unique_target(dest_dir, path.name)
        try:
            shutil.move(str(path), str(dest))
            print(f"  -> moved to {dest.resolve()}")
        except FileNotFoundError:
            print(f"  ! file not found, skipping")
        time.sleep(3)

# ---------- Trigger: watch the folder ----------
class Handler(FileSystemEventHandler):
    def on_created(self, event):
        if not event.is_directory:
            process(Path(event.src_path))
    def on_moved(self, event):
        if not event.is_directory:
            process(Path(event.dest_path))

if __name__ == "__main__":
    if not os.getenv("LLM_API_KEY") or not MODEL:
        raise SystemExit("Missing LLM_API_KEY or LLM_MODEL in .env")
    INBOX.mkdir(exist_ok=True)
    OUT.mkdir(exist_ok=True)
    observer = Observer()
    observer.schedule(Handler(), str(INBOX), recursive=False)
    observer.start()
    print(f"Watching '{INBOX}/' ... drop files in. Ctrl+C to stop.")
    try:
        while True:
            for f in list(INBOX.iterdir()):     # pick up new or retry-pending files
                process(f)
            time.sleep(15)
    except KeyboardInterrupt:
        observer.stop()
    observer.join()
